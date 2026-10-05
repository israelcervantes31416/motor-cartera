"""Leer una fuente oficial por lotes, con la fila y la hoja de cada registro.

cartera/v1 lee su archivo entero en un DataFrame. Una fuente oficial puede ser de 500,000 filas por
93 columnas: se lee por lotes (MC_FILAS_POR_LOTE filas) y nunca esta entera en memoria.

- csv: por bloques, con pandas, sobre el texto en la codificacion detectada al recorrerlo antes.
- zip: los csv de adentro, por bloques, descomprimiendo mientras se leen.
- xlsx: con openpyxl en modo de solo lectura, que recorre la hoja como un flujo.

La estructura se revisa antes de leer un solo registro: el encabezado tiene que ser exactamente
el del contrato (ver `contratos.fuente.verificar_encabezado`). El encabezado es la fila 1; cada
registro conserva su fila en el archivo y su hoja: la de un xlsx, o el miembro de un zip. Todo se
entrega como texto sin espacios alrededor, con un vacio como NA: convertir es trabajo del contrato.
Una fila completamente vacia no es un registro, pero sigue contando para numerar las demas.

Una fuente puede traer una hoja companera (CARRIER, junto a CARTERA). No se juzga ni se publica: se
reconoce y se audita sin bloquear nada, salvo que el archivo este danado (ver `auditar_companera`).
"""

from __future__ import annotations

import codecs
import csv
import io
import re
import zipfile
import zlib
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path, PurePosixPath
from typing import TextIO

import pandas as pd

from motor_cartera.contratos.fuente import (
    ContratoFuente,
    ErrorDeEstructura,
    normalizar_encabezado,
    verificar_encabezado,
)
from motor_cartera.fuentes.formatos import BLOQUE, Formato, FormatoNoCorresponde, verificar_texto
from motor_cartera.ingesta.lectores import ErrorDeLectura

PRIMERA_FILA_DE_DATOS = 2


@dataclass(frozen=True)
class Lote:
    """Unas filas de la tabla principal, con las columnas del contrato en su orden oficial."""

    datos: pd.DataFrame
    """Texto sin espacios alrededor, un vacio como NA. El indice es la fila del archivo."""
    hoja: str | None
    """La hoja del xlsx o el miembro del zip de donde salieron; None en un csv suelto."""


@dataclass(frozen=True)
class EspecificacionCompanera:
    """Que forma se espera de una hoja companera y que se audita de ella."""

    nombre: str
    """Como se llama la hoja en un xlsx; en un zip, lo que contiene el nombre del miembro."""
    columnas: tuple[str, ...]
    llave: str
    """La columna que tiene que existir en la tabla principal: una companera no trae registros de
    otra cosa."""
    formas: dict[str, str] = field(default_factory=dict)
    """Columnas con una forma exacta, como el telefono: se cuentan las filas que no la tienen."""


@dataclass(frozen=True)
class CompaneraAuditada:
    """Lo que se encontro de una hoja companera."""

    nombre: str
    filas: int
    columnas: int
    estructura_reconocida: bool
    """Si trae exactamente las columnas esperadas."""
    advertencias: tuple[str, ...]
    """Lo incoherente, para auditarlo. No bloquea la publicacion de la tabla principal."""


def abrir_fuente(
    ruta: Path,
    formato: Formato,
    nombre: str,
    contrato: ContratoFuente,
    *,
    filas_por_lote: int,
    hoja_principal: str | None = None,
    companera: EspecificacionCompanera | None = None,
) -> Fuente:
    """La fuente lista para leerse por lotes, con su estructura ya revisada.

    Levanta ErrorDeLectura si el archivo no se puede leer o no se sabe que tabla es la principal, y
    ErrorDeEstructura si la tabla principal no tiene exactamente las columnas del contrato. En un
    xlsx, la tabla principal es la hoja `hoja_principal` (sin distinguir mayusculas); sin ella, la
    unica hoja con la estructura del contrato. En un zip, el unico csv con esa estructura.
    """
    if filas_por_lote < 1:
        raise ValueError("Se lee de a una fila al menos.")
    lectores = {Formato.CSV: _FuenteCsv, Formato.ZIP: _FuenteZip, Formato.XLSX: _FuenteXlsx}
    if formato not in lectores:
        raise ErrorDeLectura(f"{nombre!r}: una fuente oficial no se lee de un {formato}.")
    return lectores[formato](ruta, nombre, contrato, filas_por_lote, hoja_principal, companera)


class Fuente:
    """Una fuente abierta. Se cierra con `with` o con close()."""

    origen: str
    """Que se leyo y por que: la hoja o el miembro elegido, y lo que se descarto."""

    def __init__(
        self,
        ruta: Path,
        nombre: str,
        contrato: ContratoFuente,
        filas_por_lote: int,
        hoja_principal: str | None,
        companera: EspecificacionCompanera | None,
    ) -> None:
        self._ruta = ruta
        self._nombre = nombre
        self._contrato = contrato
        self._filas_por_lote = filas_por_lote
        self._hoja_principal = hoja_principal
        self._companera = companera
        self.origen = repr(nombre)

    def lotes(self) -> Iterator[Lote]:
        raise NotImplementedError

    def auditar_companera(self, claves: set[str]) -> CompaneraAuditada | None:
        """Audita la hoja companera, si hay una: su estructura, cuantas filas trae, cuantas son de
        una llave que la tabla principal no tiene (`claves`) y cuantas no tienen la forma esperada.
        Lo que encuentre son advertencias. Solo un archivo danado, que no se puede terminar de leer,
        levanta ErrorDeEstructura."""
        return None

    def close(self) -> None:
        pass

    def __enter__(self) -> Fuente:
        return self

    def __exit__(self, *_: object) -> None:
        self.close()


# --- csv ------------------------------------------------------------------------------------------


class _FuenteCsv(Fuente):
    def __init__(self, *argumentos) -> None:
        super().__init__(*argumentos)
        try:
            self._codificacion = verificar_texto(self._ruta, self._nombre)
        except FormatoNoCorresponde as exc:
            raise ErrorDeLectura(str(exc)) from exc
        with self._ruta.open(encoding=self._codificacion, newline="") as texto:
            self._encabezado = _verificar(
                _encabezado_csv(texto), self._contrato, f"{self._nombre!r}"
            )
        nota = "" if self._codificacion == "utf-8-sig" else f" (codificacion {self._codificacion})"
        self.origen = f"{self._nombre!r}{nota}"

    def lotes(self) -> Iterator[Lote]:
        with self._ruta.open(encoding=self._codificacion, newline="") as texto:
            texto.readline()
            for bloque in _bloques_csv(
                texto, self._encabezado, self._contrato, self._filas_por_lote
            ):
                yield Lote(bloque, None)


# --- zip ------------------------------------------------------------------------------------------


class _FuenteZip(Fuente):
    def __init__(self, *argumentos) -> None:
        super().__init__(*argumentos)
        try:
            self._paquete = zipfile.ZipFile(self._ruta)
        except (zipfile.BadZipFile, OSError) as exc:
            raise ErrorDeLectura(f"{self._nombre!r} no es un zip legible: {exc}") from exc
        principales: list[tuple[zipfile.ZipInfo, list[str]]] = []
        self._companeras: list[zipfile.ZipInfo] = []
        descartes: list[str] = []
        for info in self._paquete.infolist():
            interno = PurePosixPath(info.filename)
            if info.is_dir() or "__MACOSX" in interno.parts:
                continue
            if self._companera and self._companera.nombre in interno.stem.upper():
                self._companeras.append(info)
                continue
            if interno.suffix.lower() != ".csv":
                descartes.append(f"{info.filename!r}: dentro de un zip solo se leen csv")
                continue
            try:
                encabezado = _verificar(
                    _encabezado_de_miembro(self._paquete, info),
                    self._contrato,
                    repr(info.filename),
                )
            except (ErrorDeEstructura, ErrorDeLectura) as exc:
                descartes.append(f"{info.filename!r}: {exc}")
            else:
                principales.append((info, encabezado))
        if len(principales) != 1:
            self._paquete.close()
            if principales:
                nombres = ", ".join(repr(p[0].filename) for p in principales)
                raise ErrorDeLectura(
                    f"Hay {len(principales)} archivos con la estructura de "
                    f"{self._contrato.version} en {self._nombre!r} ({nombres}); no se elige uno a "
                    "ciegas."
                )
            motivos = "; ".join(descartes) or "no contiene nada"
            raise ErrorDeLectura(
                f"No hay un archivo con la estructura de {self._contrato.version} en "
                f"{self._nombre!r}: {motivos}"
            )
        self._miembro, self._encabezado = principales[0]
        # Solo del elegido se recorre el contenido entero: tiene que ser texto de principio a fin.
        try:
            self._codificacion = _codificacion_de_miembro(self._paquete, self._miembro)
        except ErrorDeLectura:
            self._paquete.close()
            raise
        elegido = f"{self._miembro.filename!r}, dentro de {self._nombre!r}"
        self.origen = _describir(elegido, descartes)

    def lotes(self) -> Iterator[Lote]:
        with _texto_de_miembro(self._paquete, self._miembro, self._codificacion) as texto:
            texto.readline()
            for bloque in _bloques_csv(
                texto, self._encabezado, self._contrato, self._filas_por_lote
            ):
                yield Lote(bloque, self._miembro.filename)

    def auditar_companera(self, claves: set[str]) -> CompaneraAuditada | None:
        if not self._companera or not self._companeras:
            return None
        if len(self._companeras) > 1:
            nombres = ", ".join(repr(i.filename) for i in self._companeras)
            return CompaneraAuditada(
                self._companera.nombre,
                0,
                0,
                False,
                (
                    f"Hay {len(self._companeras)} miembros {self._companera.nombre} "
                    f"({nombres}); no se audito ninguno.",
                ),
            )
        info = self._companeras[0]
        try:
            codificacion = _codificacion_de_miembro(self._paquete, info)
            with _texto_de_miembro(self._paquete, info, codificacion) as texto:
                encabezado = _encabezado_csv(texto)
                return _auditar(
                    info.filename, encabezado, _filas_csv(texto), self._companera, claves
                )
        except ErrorDeLectura as exc:
            return CompaneraAuditada(info.filename, 0, 0, False, (str(exc),))
        except (zipfile.BadZipFile, zlib.error, EOFError) as exc:
            raise ErrorDeEstructura(
                f"{info.filename!r}, dentro de {self._nombre!r}, esta danado: {exc}"
            ) from exc

    def close(self) -> None:
        self._paquete.close()


# --- xlsx -----------------------------------------------------------------------------------------


class _FuenteXlsx(Fuente):
    def __init__(self, *argumentos) -> None:
        super().__init__(*argumentos)
        import openpyxl

        # Se abre el archivo y se le pasa a openpyxl: el objeto del almacen se llama por su
        # SHA-256, sin extension, y openpyxl no abre una ruta que no termine en .xlsx. En modo de
        # solo lectura, el archivo tiene que seguir abierto mientras se lee la hoja.
        self._archivo = self._ruta.open("rb")
        try:
            self._libro = openpyxl.load_workbook(self._archivo, read_only=True, data_only=True)
        except Exception as exc:  # openpyxl tiene muchas formas de decir "esto no es un Excel"
            self._archivo.close()
            raise ErrorDeLectura(f"{self._nombre!r} no es un Excel legible: {exc}") from exc
        try:
            self._hoja, self._encabezado, descartes = self._elegir_hoja()
        except Exception:
            self.close()
            raise
        self.origen = _describir(f"hoja {self._hoja!r} de {self._nombre!r}", descartes)

    def _elegir_hoja(self) -> tuple[str, list[str], list[str]]:
        hojas = self._libro.sheetnames
        if self._hoja_principal is not None:
            iguales = [h for h in hojas if h.strip().upper() == self._hoja_principal.upper()]
            if len(iguales) != 1:
                cuantas = "no tiene la hoja" if not iguales else "tiene varias hojas que se llaman"
                raise ErrorDeLectura(
                    f"{self._nombre!r} {cuantas} {self._hoja_principal}. Sus hojas: "
                    f"{', '.join(repr(h) for h in hojas)}."
                )
            encabezado = _primera_fila(self._libro[iguales[0]])
            return (
                iguales[0],
                _verificar(encabezado, self._contrato, f"La hoja {iguales[0]!r}"),
                [],
            )
        candidatas, descartes = [], []
        for hoja in hojas:
            try:
                encabezado = _verificar(
                    _primera_fila(self._libro[hoja]), self._contrato, f"La hoja {hoja!r}"
                )
            except ErrorDeEstructura as exc:
                descartes.append(f"hoja {hoja!r}: {exc}")
            else:
                candidatas.append((hoja, encabezado))
        if len(candidatas) != 1:
            if candidatas:
                raise ErrorDeLectura(
                    f"Hay {len(candidatas)} hojas con la estructura de {self._contrato.version} "
                    f"en {self._nombre!r}; no se elige una a ciegas."
                )
            raise ErrorDeLectura(
                f"No hay una hoja con la estructura de {self._contrato.version} en "
                f"{self._nombre!r}: {'; '.join(descartes)}"
            )
        return candidatas[0][0], candidatas[0][1], descartes

    def lotes(self) -> Iterator[Lote]:
        hoja = self._libro[self._hoja]
        ancho = len(self._encabezado)
        filas: list[list[str | None]] = []
        numeros: list[int] = []
        for numero, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
            if not fila:
                continue
            if len(fila) > ancho and any((_texto_de_celda(v) or "").strip() for v in fila[ancho:]):
                raise ErrorDeEstructura(
                    f"La fila {numero} de la hoja {self._hoja!r} trae valores fuera de las {ancho} "
                    "columnas del encabezado."
                )
            valores = [_texto_de_celda(v) for v in fila[:ancho]]
            valores.extend([None] * (ancho - len(valores)))
            filas.append(valores)
            numeros.append(numero)
            if len(filas) >= self._filas_por_lote:
                yield Lote(_bloque(filas, numeros, self._encabezado, self._contrato), self._hoja)
                filas, numeros = [], []
        if filas:
            yield Lote(_bloque(filas, numeros, self._encabezado, self._contrato), self._hoja)

    def auditar_companera(self, claves: set[str]) -> CompaneraAuditada | None:
        if not self._companera:
            return None
        esperada = self._companera.nombre.upper()
        iguales = [h for h in self._libro.sheetnames if h.strip().upper() == esperada]
        if not iguales:
            return None
        if len(iguales) > 1:
            return CompaneraAuditada(
                self._companera.nombre,
                0,
                0,
                False,
                (
                    f"El libro tiene {len(iguales)} hojas que se llaman "
                    f"{self._companera.nombre}; no se audito ninguna.",
                ),
            )
        hoja = self._libro[iguales[0]]
        try:
            encabezado = _primera_fila(hoja)
            filas = (
                [_texto_de_celda(v) for v in fila]
                for fila in hoja.iter_rows(min_row=2, values_only=True)
                if fila
            )
            return _auditar(iguales[0], encabezado, filas, self._companera, claves)
        except ErrorDeEstructura:
            raise
        except Exception as exc:
            raise ErrorDeEstructura(
                f"La hoja {iguales[0]!r} de {self._nombre!r} esta danada: {exc}"
            ) from exc

    def close(self) -> None:
        self._libro.close()
        self._archivo.close()


# --- lo comun -------------------------------------------------------------------------------------


def _verificar(encabezado: list[str], contrato: ContratoFuente, donde: str) -> list[str]:
    """El encabezado revisado contra el contrato, sin las celdas vacias del final, que un xlsx o un
    csv pueden traer de mas sin que sean columnas."""
    while encabezado and not normalizar_encabezado(encabezado[-1]):
        encabezado = encabezado[:-1]
    return verificar_encabezado(encabezado, contrato, donde)


def _encabezado_csv(texto: TextIO) -> list[str]:
    linea = texto.readline()
    if not linea:
        raise ErrorDeLectura("El archivo no trae ni siquiera el encabezado.")
    return next(csv.reader([linea]))


def _bloques_csv(
    texto: TextIO, encabezado: list[str], contrato: ContratoFuente, filas_por_lote: int
) -> Iterator[pd.DataFrame]:
    """El resto del csv, despues del encabezado, por bloques. Una fila con mas campos que el
    encabezado es un error de estructura: no se adivina a que columna iban."""
    try:
        lector = pd.read_csv(
            texto,
            header=None,
            names=list(range(len(encabezado))),
            dtype="string",
            keep_default_na=False,
            na_filter=False,
            skip_blank_lines=False,
            chunksize=filas_por_lote,
            engine="c",
        )
        for bloque in lector:
            bloque.columns = encabezado
            bloque.index = bloque.index + PRIMERA_FILA_DE_DATOS
            limpio = _limpiar(bloque[list(contrato.nombres)])
            if not limpio.empty:
                yield limpio
    except pd.errors.EmptyDataError:
        return
    except pd.errors.ParserError as exc:
        raise ErrorDeEstructura(f"El csv no tiene la forma de una tabla: {exc}") from exc


def _filas_csv(texto: TextIO) -> Iterator[list[str | None]]:
    """Las filas de un csv companero, una por una. Una fila con otro numero de campos se cuenta
    igual: aqui se audita, no se juzga."""
    for fila in csv.reader(texto):
        if fila:
            yield list(fila)


def _bloque(
    filas: list[list[str | None]], numeros: list[int], encabezado: list[str], contrato
) -> pd.DataFrame:
    datos = pd.DataFrame(filas, columns=encabezado, index=numeros, dtype="string")
    return _limpiar(datos[list(contrato.nombres)])


def _limpiar(datos: pd.DataFrame) -> pd.DataFrame:
    """Sin espacios alrededor de cada valor; una celda vacia o solo con espacios es NA; una fila
    completamente vacia no es un registro."""
    limpio = {}
    for columna in datos.columns:
        valores = datos[columna].astype("string").str.strip()
        limpio[columna] = valores.mask(valores == "")
    resultado = pd.DataFrame(limpio, index=datos.index)
    return resultado[resultado.notna().any(axis=1).to_numpy()]


def _primera_fila(hoja) -> list[str]:
    for fila in hoja.iter_rows(min_row=1, max_row=1, values_only=True):
        return ["" if v is None else str(v) for v in fila]
    return []


def _texto_de_celda(valor: object) -> str | None:
    """El valor de una celda de Excel como texto, sin perder nada de lo que dice.

    Un numero entero sale sin decimales y uno con decimales en su representacion mas corta que se
    lee igual (0.1 y no 0.1000000000000000055); una fecha a medianoche como AAAA-MM-DD, y una con
    hora como AAAA-MM-DDTHH:MM:SS. Un numero que no es exacto no se redondea: el contrato decide.
    """
    if valor is None:
        return None
    if isinstance(valor, str):
        return valor
    if isinstance(valor, bool):
        return "TRUE" if valor else "FALSE"
    if isinstance(valor, int):
        return str(valor)
    if isinstance(valor, float):
        if valor.is_integer() and abs(valor) < 2**53:
            return str(int(valor))
        return repr(valor)
    if isinstance(valor, datetime):
        if valor.time() == time(0) and valor.tzinfo is None:
            return valor.date().isoformat()
        return valor.isoformat(sep="T", timespec="microseconds" if valor.microsecond else "seconds")
    if isinstance(valor, date | time):
        return valor.isoformat()
    if isinstance(valor, Decimal):
        return format(valor, "f")
    return str(valor)


def _codificacion_de_miembro(paquete: zipfile.ZipFile, info: zipfile.ZipInfo) -> str:
    """Recorre el miembro por bloques, como verificar_texto con un archivo, y devuelve su
    codificacion; ErrorDeLectura si no es texto."""
    for codificacion in ("utf-8-sig", "cp1252"):
        decodificador = codecs.getincrementaldecoder(codificacion)()
        try:
            with paquete.open(info) as miembro:
                while bloque := miembro.read(BLOQUE):
                    if b"\x00" in bloque:
                        raise ErrorDeLectura(f"{info.filename!r} trae bytes nulos: no es texto.")
                    decodificador.decode(bloque)
                decodificador.decode(b"", final=True)
        except UnicodeDecodeError:
            continue
        return codificacion
    raise ErrorDeLectura(f"{info.filename!r} no es texto en UTF-8 ni en cp1252.")


def _texto_de_miembro(paquete: zipfile.ZipFile, info: zipfile.ZipInfo, codificacion: str):
    return io.TextIOWrapper(paquete.open(info), encoding=codificacion, newline="")


def _encabezado_de_miembro(paquete: zipfile.ZipFile, info: zipfile.ZipInfo) -> list[str]:
    """El encabezado de un miembro, leyendo solo su primera linea: alcanza para saber que tabla es.
    Su contenido completo se revisa despues, si es el elegido."""
    with paquete.open(info) as miembro:
        linea = miembro.readline()
    for codificacion in ("utf-8-sig", "cp1252"):
        try:
            texto = linea.decode(codificacion)
        except UnicodeDecodeError:
            continue
        if not texto.strip():
            raise ErrorDeLectura("no trae ni siquiera el encabezado")
        return next(csv.reader([texto]))
    raise ErrorDeLectura("su encabezado no es texto en UTF-8 ni en cp1252")


def _auditar(
    nombre: str,
    encabezado: list[str],
    filas: Iterator[list[str | None]],
    esperada: EspecificacionCompanera,
    claves: set[str],
) -> CompaneraAuditada:
    """Lo que se puede decir de una companera sin juzgarla registro por registro."""
    nombres = [normalizar_encabezado(h) for h in encabezado]
    while nombres and not nombres[-1]:
        nombres.pop()
    advertencias: list[str] = []
    repetidos = sorted({n for n in nombres if n and nombres.count(n) > 1})
    faltantes = [c for c in esperada.columnas if c not in nombres]
    sobrantes = [n for n in dict.fromkeys(nombres) if n and n not in esperada.columnas]
    if repetidos:
        advertencias.append(f"Encabezados repetidos: {', '.join(repetidos)}.")
    if faltantes:
        advertencias.append(f"Le faltan {len(faltantes)} columnas: {', '.join(faltantes)}.")
    if sobrantes:
        advertencias.append(f"Le sobran {len(sobrantes)} columnas: {', '.join(sobrantes)}.")
    reconocida = not (repetidos or faltantes or sobrantes or "" in nombres)

    posicion = {n: i for i, n in reversed(list(enumerate(nombres)))}
    llave = posicion.get(esperada.llave)
    formas = {c: (posicion[c], re.compile(p)) for c, p in esperada.formas.items() if c in posicion}
    total = ajenas = 0
    sin_forma = dict.fromkeys(formas, 0)
    for fila in filas:
        valores = [None if v is None else v.strip() for v in fila]
        if not any(valores):
            continue
        total += 1
        if llave is not None and (len(valores) <= llave or valores[llave] not in claves):
            ajenas += 1
        for columna, (i, patron) in formas.items():
            valor = valores[i] if i < len(valores) else None
            if not valor or not patron.fullmatch(valor):
                sin_forma[columna] += 1
    if total == 0:
        advertencias.append("No trae registros.")
    if ajenas:
        advertencias.append(
            f"{ajenas:,} filas con un {esperada.llave} que no esta en la tabla principal."
        )
    for columna, cuantas in sin_forma.items():
        if cuantas:
            advertencias.append(
                f"{cuantas:,} filas con un {columna} que no tiene la forma esperada."
            )
    return CompaneraAuditada(nombre, total, len(nombres), reconocida, tuple(advertencias))


def _describir(elegido: str, descartes: list[str]) -> str:
    if not descartes:
        return elegido
    return f"{elegido}. Descartado: {'; '.join(descartes)}"
