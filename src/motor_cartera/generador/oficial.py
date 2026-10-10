"""Generador de las fuentes oficiales sinteticas: CARTERA (93 columnas), CARRIER y PAGOS.

REGLA DURA DEL PROYECTO: ningun dato real entra a este repositorio. Nada de aqui sale de un archivo
real: ni nombres, ni domicilios, ni telefonos, ni clientes, ni coordenadas, ni distribuciones de un
archivo privado. Lo que se conserva son rasgos genericos de cualquier cartera: saldos con cola
larga, moras distintas, concentracion territorial (con la poblacion del censo del INEGI, que es
publica), varios telefonos por cliente, clientes con y sin pagos.

Algunas decisiones para que lo sintetico no se confunda con algo real:

- Los telefonos empiezan con 0: ningun numero nacional de Mexico empieza asi, asi que no se le puede
  marcar a nadie. La CLAVE_SPEI empieza con 000, que no es la clave de ningun banco.
- Los nombres y las calles salen de listas genericas: no corresponden a ninguna persona.
- LATITUD y LONGITUD van vacias. El catalogo no trae centroides municipales, y una coordenada al
  azar seria incoherente con la poblacion del cliente. La geografia con coordenadas es posterior.

Como escala. Cada cliente se describe con un estado compacto (arreglos de numpy) y sus atributos
fijos (nombre, domicilio, telefonos, producto, municipio) se derivan de un hash de su identificador
y la semilla: son los mismos sin importar en que lote, en que orden o en que corte se generen. Las
93 columnas se arman por bloques con pyarrow y se escriben por bloques: 1,000,000 de cuentas nunca
estan enteras en memoria como texto.
"""

from __future__ import annotations

import hashlib
import json
import zipfile
from collections.abc import Iterator
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.compute as pc
import pyarrow.csv as pcsv

from motor_cartera.contratos.cartera import CANALES, PRODUCTOS
from motor_cartera.contratos.cartera_v2 import (
    COLUMNAS_CARRIER,
    CONTRATO_V2,
    HOJA_CARRIER,
    HOJA_CARTERA,
)
from motor_cartera.contratos.fuente import ContratoFuente, Tipo
from motor_cartera.contratos.pagos import CONTRATO_PAGOS
from motor_cartera.fuentes.geografia import catalogo, normalizar_nombre
from motor_cartera.segmentacion import TRAMOS_ATRASO

PERFILES = {
    "XS": 1_000,
    "S": 10_000,
    "M": 100_000,
    "L": 250_000,
    "XL": 500_000,
    "XXL": 1_000_000,
}
"""Cuantas cuentas trae cada perfil. XL es el escenario empresarial objetivo; XXL, la prueba de
esfuerzo. Ninguno corre en el CI: el CI usa XS y menos."""

FILAS_POR_BLOQUE = 50_000
"""De cuantas cuentas se arma y se escribe cada bloque."""

# --- distribuciones genericas --------------------------------------------------------------------

PESO_ENTIDAD = {"21": 0.70, "29": 0.08, "30": 0.08, "15": 0.08, "09": 0.06}
"""Casi toda la cartera en una entidad, como una cartera regional: las cinco de cartera/v1. Dentro
de cada entidad, los municipios pesan por su poblacion del censo."""

PESO_PRODUCTO = {"CONSUMO": 0.45, "TARJETA": 0.30, "NOMINA": 0.15, "AUTOMOTRIZ": 0.10}
MEDIANA_SALDO = {"CONSUMO": 15_000, "TARJETA": 12_000, "NOMINA": 20_000, "AUTOMOTRIZ": 90_000}
PESO_TRAMO = {"0": 0.10, "1-30": 0.30, "31-60": 0.20, "61-90": 0.15, "91+": 0.25}
TOPE_ATRASO = 720
CANAL_POR_TRAMO = {
    "0": {"DIGITAL": 0.60, "TELEFONICA": 0.35, "CAMPO": 0.05},
    "1-30": {"DIGITAL": 0.45, "TELEFONICA": 0.45, "CAMPO": 0.10},
    "31-60": {"DIGITAL": 0.25, "TELEFONICA": 0.55, "CAMPO": 0.20},
    "61-90": {"DIGITAL": 0.10, "TELEFONICA": 0.50, "CAMPO": 0.40},
    "91+": {"DIGITAL": 0.05, "TELEFONICA": 0.35, "CAMPO": 0.60},
}
ESTRATEGIA_POR_TRAMO = {
    "0": "PREVENTIVA",
    "1-30": "TEMPRANA",
    "31-60": "INTERMEDIA",
    "61-90": "TARDIA",
    "91+": "RECUPERACION",
}
CODIGOS_POSTALES = {
    "09": (1000, 16999),
    "15": (50000, 57999),
    "21": (72000, 75999),
    "29": (90000, 90999),
    "30": (91000, 96999),
}
"""Los intervalos publicos de codigos postales de cada entidad. El codigo es sintetico: no se
valida contra el municipio."""

NOMBRES = (
    "ANA", "LUIS", "MARIA", "JOSE", "ROSA", "JUAN", "LAURA", "CARLOS", "ELENA", "MIGUEL",
    "SOFIA", "JORGE", "PATRICIA", "PEDRO", "CLAUDIA", "RAUL", "MONICA", "ARTURO", "GABRIELA",
    "RICARDO", "VERONICA", "FERNANDO", "ADRIANA", "ALEJANDRO", "TERESA", "ROBERTO", "LUCIA",
    "SERGIO", "MARTHA", "ENRIQUE", "SILVIA", "JAVIER", "ISABEL", "DANIEL", "NORMA", "OSCAR",
)  # fmt: skip
APELLIDOS = (
    "HERNANDEZ", "GARCIA", "MARTINEZ", "LOPEZ", "GONZALEZ", "PEREZ", "RODRIGUEZ", "SANCHEZ",
    "RAMIREZ", "CRUZ", "FLORES", "GOMEZ", "MORALES", "VAZQUEZ", "JIMENEZ", "REYES", "DIAZ",
    "TORRES", "GUTIERREZ", "RUIZ", "MENDOZA", "AGUILAR", "ORTIZ", "MORENO", "CASTILLO",
    "ROMERO", "ALVAREZ", "MENDEZ", "CHAVEZ", "RIVERA", "JUAREZ", "RAMOS", "DOMINGUEZ",
)  # fmt: skip
CALLES = (
    "HIDALGO", "JUAREZ", "MORELOS", "INDEPENDENCIA", "REFORMA", "5 DE MAYO",
    "16 DE SEPTIEMBRE", "ALLENDE", "ZARAGOZA", "GUERRERO", "MATAMOROS", "ALDAMA", "NIÑOS HEROES",
    "LIBERTAD", "LAS ROSAS", "DEL TRABAJO", "PRIMAVERA", "LOS PINOS",
)  # fmt: skip
COLONIAS = (
    "CENTRO", "SAN JOSE", "LA PAZ", "EL CARMEN", "SANTA MARIA", "LOS PINOS", "LAS FLORES",
    "INDUSTRIAL", "MODERNA", "JARDINES", "SAN MIGUEL", "LA LOMA", "EL PROGRESO", "AMPLIACION",
)  # fmt: skip
OCUPACIONES = (
    "EMPLEADO", "COMERCIANTE", "HOGAR", "OBRERO", "PROFESIONISTA", "CHOFER", "JUBILADO",
    "ESTUDIANTE", "AGRICULTOR", "OTRO",
)  # fmt: skip
TIPOS_TELEFONO = ("CELULAR", "CASA", "TRABAJO", "RECADOS")
REFERENCIAS = (
    "FRENTE A LA TIENDA", "JUNTO A LA ESCUELA", "CASA DE DOS PISOS", "PORTON NEGRO",
    "ENTRE CALLES", "A UN LADO DE LA IGLESIA",
)  # fmt: skip
GESTIONES = (
    "SIN CONTACTO", "PROMESA DE PAGO", "NO LOCALIZADO", "RECADO", "NEGATIVA DE PAGO",
    "CONVENIO", "TELEFONO EQUIVOCADO",
)  # fmt: skip
NOMBRE_DESPACHO = "DESPACHO SINTETICO"
LEEME = (
    "Fuente oficial sintetica generada por motor-cartera. Ningun dato corresponde a una persona: "
    "los telefonos empiezan con 0 y no se pueden marcar."
)
FECHA_EN_ZIP = (1980, 1, 1, 0, 0, 0)
"""La fecha de cada miembro de un zip generado: la minima del formato, siempre la misma. Asi el
mismo contenido produce los mismos bytes, y un escenario se verifica por su SHA-256."""

# --- el hash que fija los atributos de cada cliente -----------------------------------------------

_M1 = np.uint64(0xBF58476D1CE4E5B9)
_M2 = np.uint64(0x94D049BB133111EB)
_DORADO = np.uint64(0x9E3779B97F4A7C15)


def _mezclar(x: np.ndarray) -> np.ndarray:
    """splitmix64: de un entero, otro que parece al azar, igual en cualquier maquina."""
    with np.errstate(over="ignore"):
        x = x.astype(np.uint64) + _DORADO
        x = (x ^ (x >> np.uint64(30))) * _M1
        x = (x ^ (x >> np.uint64(27))) * _M2
        return x ^ (x >> np.uint64(31))


def _uniforme(clientes: np.ndarray, semilla: int, campo: int) -> np.ndarray:
    """Un numero en [0, 1) para cada cliente y cada atributo, fijo para esa semilla."""
    sal = _mezclar(np.array([semilla * 1_000_003 + campo], dtype=np.uint64))[0]
    with np.errstate(over="ignore"):
        bits = _mezclar(clientes.astype(np.uint64) ^ sal)
    return (bits >> np.uint64(11)).astype(np.float64) / float(1 << 53)


def _elegir(u: np.ndarray, pesos: list[float]) -> np.ndarray:
    acumulado = np.cumsum(np.asarray(pesos, dtype=np.float64))
    acumulado /= acumulado[-1]
    return np.minimum(np.searchsorted(acumulado, u, side="right"), len(pesos) - 1)


# --- la geografia sintetica -----------------------------------------------------------------------


@dataclass(frozen=True)
class _Geografia:
    """Los municipios que el generador usa, con su entidad, su nombre normalizado y su peso."""

    cve_entidad: np.ndarray
    cve_municipio: np.ndarray
    estado: np.ndarray
    poblacion: np.ndarray
    peso: np.ndarray


def _geografia() -> _Geografia:
    cat = catalogo()
    filas = []
    for cve, peso_entidad in PESO_ENTIDAD.items():
        entidad = cat.entidad(cve)
        municipios = [m for m in entidad.municipios if m.poblacion]
        total = sum(m.poblacion for m in municipios)
        for m in municipios:
            # Un nombre ambiguo no se usa: el generador produce datos que la proyeccion resuelve.
            if cat.por_municipio.get((cve, normalizar_nombre(m.nombre))) != m.cve_municipio:
                continue
            filas.append(
                (
                    cve,
                    m.cve_municipio,
                    normalizar_nombre(entidad.nombre),
                    normalizar_nombre(m.nombre),
                    peso_entidad * m.poblacion / total,
                )
            )
    columnas = list(zip(*filas, strict=True))
    return _Geografia(
        np.array(columnas[0]),
        np.array(columnas[1]),
        np.array(columnas[2]),
        np.array(columnas[3]),
        np.array(columnas[4], dtype=np.float64),
    )


# --- el estado de la cartera ----------------------------------------------------------------------


@dataclass(frozen=True)
class EstadoCartera:
    """Las cuentas de un corte, en arreglos paralelos. Lo que no esta aqui se deriva del cliente."""

    cliente: np.ndarray
    """El numero de cada cliente: CLIENTE_UNICO es CU seguido de ese numero en 10 digitos."""
    saldo_centavos: np.ndarray
    moratorios_centavos: np.ndarray
    dias_atraso: np.ndarray
    atraso_maximo: np.ndarray
    asignacion: np.ndarray
    """La fecha en que la cuenta llego al despacho (datetime64[D])."""
    pagos: np.ndarray
    """Cuantos pagos ha hecho el cliente."""
    monto_pagos_centavos: np.ndarray
    ultimo_pago: np.ndarray
    """La fecha de su ultimo pago, o NaT si nunca ha pagado."""
    ultimo_pago_centavos: np.ndarray

    def __len__(self) -> int:
        return len(self.cliente)

    def tomar(self, indices: np.ndarray) -> EstadoCartera:
        return EstadoCartera(**{c: getattr(self, c)[indices] for c in self.__dataclass_fields__})


def nuevos_clientes(cuantos: int, rng: np.random.Generator, usados: set[int] | None = None):
    """Numeros de cliente distintos entre si y de `usados`: una cuenta nueva nunca reusa la
    identidad de otra, ni de una que ya salio de la cartera."""
    usados = usados if usados is not None else set()
    elegidos: list[int] = []
    while len(elegidos) < cuantos:
        faltan = cuantos - len(elegidos)
        candidatos = rng.integers(1, 10**10, size=faltan * 2 + 10, dtype=np.int64)
        for c in np.unique(candidatos):
            c = int(c)
            if c not in usados:
                usados.add(c)
                elegidos.append(c)
                if len(elegidos) == cuantos:
                    break
    return np.array(elegidos, dtype=np.int64)


def estado_inicial(
    n: int, *, semilla: int, fecha_corte: date, usados: set[int] | None = None
) -> EstadoCartera:
    """Las `n` cuentas de un primer corte."""
    rng = np.random.default_rng([semilla, 0])
    cliente = nuevos_clientes(n, rng, usados)
    return _cuentas_nuevas(cliente, rng, semilla, fecha_corte)


def _cuentas_nuevas(
    cliente: np.ndarray, rng: np.random.Generator, semilla: int, fecha_corte: date
) -> EstadoCartera:
    n = len(cliente)
    producto = _producto(cliente, semilla)
    medianas = np.array([MEDIANA_SALDO[p] for p in PRODUCTOS])[producto]
    saldo = np.round(medianas * rng.lognormal(0.0, 1.0, size=n) * 100).astype(np.int64)
    etiquetas = [e for e, _, _ in TRAMOS_ATRASO]
    tramo = rng.choice(len(TRAMOS_ATRASO), size=n, p=[PESO_TRAMO[e] for e in etiquetas])
    dias = np.zeros(n, dtype=np.int64)
    for j, (_, desde, hasta) in enumerate(TRAMOS_ATRASO):
        en_tramo = tramo == j
        dias[en_tramo] = rng.integers(
            desde, (TOPE_ATRASO if hasta is None else hasta) + 1, size=int(en_tramo.sum())
        )
    moratorios = (saldo * 0.0015 * dias).astype(np.int64)
    corte = np.datetime64(fecha_corte, "D")
    asignacion = corte - rng.integers(0, 366, size=n).astype("timedelta64[D]")
    pagos = rng.integers(0, 31, size=n)
    pago = rng.random(n) < 0.6
    ultimo = np.where(
        pago,
        corte - rng.integers(1, 181, size=n).astype("timedelta64[D]"),
        np.datetime64("NaT", "D"),
    )
    ultimo_centavos = np.where(pago, np.maximum(saldo // rng.integers(10, 60, size=n), 1000), 0)
    return EstadoCartera(
        cliente=cliente,
        saldo_centavos=saldo,
        moratorios_centavos=moratorios,
        dias_atraso=dias,
        atraso_maximo=dias + rng.integers(0, 60, size=n),
        asignacion=asignacion.astype("datetime64[D]"),
        pagos=np.where(pago, np.maximum(pagos, 1), 0),
        monto_pagos_centavos=np.where(pago, ultimo_centavos * np.maximum(pagos, 1), 0),
        ultimo_pago=ultimo.astype("datetime64[D]"),
        ultimo_pago_centavos=ultimo_centavos.astype(np.int64),
    )


def _producto(cliente: np.ndarray, semilla: int) -> np.ndarray:
    return _elegir(_uniforme(cliente, semilla, 1), [PESO_PRODUCTO[p] for p in PRODUCTOS])


# --- de un estado a las 93 columnas ---------------------------------------------------------------


def _texto(valores) -> pa.Array:
    return pa.array(valores, type=pa.string())


def _de_lista(lista: tuple[str, ...], indices: np.ndarray) -> pa.Array:
    return _texto(list(lista)).take(pa.array(indices, type=pa.int64()))


def _o_nulo(arreglo: pa.Array, presente: np.ndarray) -> pa.Array:
    return pc.if_else(pa.array(presente), arreglo, pa.nulls(len(arreglo), pa.string()))


def _numero(enteros: np.ndarray) -> pa.Array:
    return pc.cast(pa.array(np.asarray(enteros, dtype=np.int64)), pa.string())


def _digitos(enteros: np.ndarray, ancho: int) -> pa.Array:
    return pc.utf8_lpad(_numero(enteros), ancho, "0")


def _pesos(centavos: np.ndarray) -> pa.Array:
    """Centavos a texto con punto decimal y dos decimales: 150050 -> 1500.50."""
    centavos = np.asarray(centavos, dtype=np.int64)
    signo = pc.if_else(pa.array(centavos < 0), "-", "")
    absolutos = np.abs(centavos)
    entero = _numero(absolutos // 100)
    fraccion = _digitos(absolutos % 100, 2)
    return pc.binary_join_element_wise(
        signo, pc.binary_join_element_wise(entero, fraccion, "."), ""
    )


def _fechas(dias: np.ndarray) -> pa.Array:
    return pc.cast(pa.array(dias.astype("datetime64[D]")), pa.string())


class _Azar:
    """Un numero en [0, 1) por cliente para cada atributo: cada atributo con su propio campo, para
    que dos atributos no salgan correlacionados por compartir el mismo azar."""

    def __init__(self, clientes: np.ndarray, semilla: int) -> None:
        self._clientes = clientes
        self._semilla = semilla
        self._vistos: dict[str, np.ndarray] = {}

    def __call__(self, atributo: str) -> np.ndarray:
        if atributo not in self._vistos:
            huella = hashlib.blake2b(atributo.encode("ascii"), digest_size=8).digest()
            campo = int.from_bytes(huella, "little") % 2**62
            self._vistos[atributo] = _uniforme(self._clientes, self._semilla, campo)
        return self._vistos[atributo]

    def indice(self, atributo: str, opciones: int) -> np.ndarray:
        return np.minimum((self(atributo) * opciones).astype(np.int64), opciones - 1)

    def entero(self, atributo: str, desde: int, hasta: int) -> np.ndarray:
        """Un entero en [desde, hasta]."""
        return desde + self.indice(atributo, hasta - desde + 1)


def tabla_cartera(estado: EstadoCartera, *, semilla: int, fecha_corte: date) -> pa.Table:
    """Las 93 columnas de CARTERA, como texto y en el orden oficial, para las cuentas de `estado`.

    Lo fijo de cada cliente (su nombre, su domicilio, sus telefonos, su producto, su municipio)
    sale del hash de su numero y la semilla; lo que cambia de un corte a otro (saldo, atraso,
    pagos), de `estado`.
    """
    n = len(estado)
    c = estado.cliente
    azar = _Azar(c, semilla)
    geo = _geografia()
    corte = np.datetime64(fecha_corte, "D")

    municipio = _elegir(azar("municipio"), list(geo.peso))
    cve_ent = geo.cve_entidad[municipio]
    poblacion = _texto(list(geo.poblacion[municipio]))
    estado_cte = _texto(list(geo.estado[municipio]))
    producto = _producto(c, semilla)
    dias = estado.dias_atraso
    etiquetas = [e for e, _, _ in TRAMOS_ATRASO]
    tramo = np.zeros(n, dtype=np.int64)
    for j, (_, desde, hasta) in enumerate(TRAMOS_ATRASO):
        tramo[(dias >= desde) & (dias <= (10**9 if hasta is None else hasta))] = j
    canal = np.zeros(n, dtype=np.int64)
    for j, etiqueta in enumerate(etiquetas):
        en_tramo = tramo == j
        pesos = [CANAL_POR_TRAMO[etiqueta][k] for k in CANALES]
        canal[en_tramo] = _elegir(azar("canal")[en_tramo], pesos)
    rango_cp = np.array([CODIGOS_POSTALES[e] for e in cve_ent]).reshape(-1, 2)
    codigo_postal = rango_cp[:, 0] + (azar("cp") * (rango_cp[:, 1] - rango_cp[:, 0] + 1)).astype(
        np.int64
    )

    saldo = estado.saldo_centavos
    moratorios = estado.moratorios_centavos
    total = saldo + moratorios
    semanas = -(-dias // 7)
    pago_normal = np.maximum(saldo // 52, 5000)
    atrasado = np.minimum(total, pago_normal * semanas)
    telefonos = 1 + _elegir(azar("telefonos"), [0.4, 0.3, 0.2, 0.1])
    con_aval = azar("aval") < 0.30
    con_empleo = azar("empleo") < 0.45
    con_plan = azar("plan") < 0.15
    con_promesa = azar("promesa") < 0.10
    gestionada = azar("gestionada") < 0.70
    pago = ~np.isnat(estado.ultimo_pago)

    def telefono(atributo: str) -> pa.Array:
        nueve = (azar(atributo) * 10**9).astype(np.int64)
        return pc.binary_join_element_wise("0", _digitos(nueve, 9), "")

    def nombre(prefijo: str) -> pa.Array:
        return pc.binary_join_element_wise(
            _de_lista(NOMBRES, azar.indice(f"{prefijo}n", len(NOMBRES))),
            _de_lista(APELLIDOS, azar.indice(f"{prefijo}p", len(APELLIDOS))),
            _de_lista(APELLIDOS, azar.indice(f"{prefijo}m", len(APELLIDOS))),
            " ",
        )

    def calle(prefijo: str) -> pa.Array:
        return pc.binary_join_element_wise(
            "CALLE", _de_lista(CALLES, azar.indice(f"{prefijo}c", len(CALLES))), " "
        )

    def colonia(prefijo: str) -> pa.Array:
        return pc.binary_join_element_wise(
            "COL", _de_lista(COLONIAS, azar.indice(f"{prefijo}k", len(COLONIAS))), " "
        )

    def codigo(prefijo: str, atributo: str, opciones: int, ancho: int) -> pa.Array:
        return pc.binary_join_element_wise(
            prefijo, _digitos(azar.entero(atributo, 1, opciones), ancho), " "
        )

    plazo = azar.entero("plazo", 4, 52)
    monto_plan = (total * (0.5 + 0.4 * azar("monto_plan"))).astype(np.int64)
    enganche = monto_plan // 10
    abono = np.maximum((monto_plan - enganche) // plazo, 1000)
    recibidos = (azar("recibidos") * plazo).astype(np.int64)
    ultima_gestion = corte - azar.entero("ultima_gestion", 0, 30).astype("timedelta64[D]")
    nulos = pa.nulls(n, pa.string())
    gestion = _de_lista(GESTIONES, azar.indice("gestion", len(GESTIONES)))

    columnas: dict[str, pa.Array] = {
        "CLIENTE_UNICO": pc.binary_join_element_wise("CU", _digitos(c, 10), ""),
        "NOMBRE_CTE": nombre("cliente"),
        "GENERO_CLIENTE": _de_lista(("H", "M"), azar.indice("genero", 2)),
        "EDAD_CLIENTE": _numero(azar.entero("edad", 18, 79)),
        "OCUPACION": _de_lista(OCUPACIONES, azar.indice("ocupacion", len(OCUPACIONES))),
        "DIRECCION_CTE": calle("domicilio"),
        "NUM_EXT_CTE": _o_nulo(_numero(azar.entero("numext", 1, 999)), azar("sn") < 0.92),
        "NUM_INT_CTE": _o_nulo(
            _de_lista(("A", "B", "1", "2"), azar.indice("numint", 4)), azar("connumint") < 0.15
        ),
        "CP_CTE": _digitos(codigo_postal, 5),
        "COLONIA_CTE": colonia("domicilio"),
        "POBLACION_CTE": poblacion,
        "ESTADO_CTE": estado_cte,
        "TERRITORIO": pc.binary_join_element_wise("TERRITORIO", _texto(list(cve_ent)), " "),
        "TERRITORIAL": pc.binary_join_element_wise("TERRITORIAL", _texto(list(cve_ent)), " "),
        "ZONA": codigo("ZONA", "zona", 12, 2),
        "ZONAL": codigo("ZONAL", "zonal", 4, 2),
        "NOMBRE_DESPACHO": _texto([NOMBRE_DESPACHO] * n),
        "GERENCIA": codigo("GERENCIA", "gerencia", 6, 2),
        "FECHA_ASIGNACION": _fechas(estado.asignacion),
        "DIAS_ASIGNACION": _numero((corte - estado.asignacion).astype(np.int64)),
        "REFERENCIAS_DOMICILIO": _o_nulo(
            _de_lista(REFERENCIAS, azar.indice("referencia", len(REFERENCIAS))),
            azar("conreferencia") < 0.5,
        ),
        "CLASIFICACION_CTE": _de_lista(("A", "B", "C", "D"), azar.indice("clasificacion", 4)),
        "DIQUE": _de_lista(("D1", "D2", "D3", "D4"), azar.indice("dique", 4)),
        "ATRASO_MAXIMO": _numero(estado.atraso_maximo),
        "DIAS_ATRASO": _numero(dias),
        "SALDO": _pesos(saldo),
        "MORATORIOS": _pesos(moratorios),
        "SALDO_TOTAL": _pesos(total),
        "SALDO ATRASADO": _pesos(atrasado),
        "SALDO REQUERIDO": _pesos(atrasado),
        "PAGO_NORMAL": _pesos(pago_normal),
        "PRODUCTO": _de_lista(PRODUCTOS, producto),
        "ESTRATEGIA": _de_lista(tuple(ESTRATEGIA_POR_TRAMO[e] for e in etiquetas), tramo),
        "FECHA_ULTIMO_PAGO": _o_nulo(_fechas(np.where(pago, estado.ultimo_pago, corte)), pago),
        "IMP_ULTIMO_PAGO": _o_nulo(_pesos(estado.ultimo_pago_centavos), pago),
        "CALLE_EMPLEO": _o_nulo(calle("empleo"), con_empleo),
        "NUM_EXT_EMPLEO": _o_nulo(_numero(azar.entero("numextempleo", 1, 500)), con_empleo),
        "NUM_INT_EMPLEO": nulos,
        "COLONIA_EMPLEO": _o_nulo(colonia("empleo"), con_empleo),
        "POBLACION_EMPLEO": _o_nulo(poblacion, con_empleo),
        "ESTADO_EMPLEO": _o_nulo(estado_cte, con_empleo),
        "NOMBRE_AVAL": _o_nulo(nombre("aval"), con_aval),
        "TEL_AVAL": _o_nulo(telefono("telaval"), con_aval),
        "CALLE_AVAL": _o_nulo(calle("aval"), con_aval),
        "NUM_EXT_AVAL": _o_nulo(_numero(azar.entero("numextaval", 1, 300)), con_aval),
        "COLONIA_AVAL": _o_nulo(colonia("aval"), con_aval),
        "CP_AVAL": _o_nulo(_digitos(codigo_postal, 5), con_aval),
        "POBLACION_AVAL": _o_nulo(poblacion, con_aval),
        "ESTADO_AVAL": _o_nulo(estado_cte, con_aval),
        "FIDIAPAGO": _de_lista(("S", "N"), (azar("fidiapago") >= 0.2).astype(np.int64)),
        "LATITUD": nulos,
        "LONGITUD": nulos,
        "DESPACHO_GESTIONO": _o_nulo(_texto([NOMBRE_DESPACHO] * n), gestionada),
        "ULTIMA_GESTION": _o_nulo(_fechas(ultima_gestion), gestionada),
        "GESTION_DESC": _o_nulo(gestion, gestionada),
        "CAMPANIA_RELAMPAGO": _o_nulo(_texto(["REL-01"] * n), azar("relampago") < 0.05),
        "CAMPANIA": pc.binary_join_element_wise(
            "CAMP", _digitos(azar.entero("campania", 1, 6), 2), "-"
        ),
        "PREVENTA": _de_lista(("S", "N"), (azar("preventa") >= 0.1).astype(np.int64)),
        "ID_GRUPO": _o_nulo(
            pc.binary_join_element_wise("G", _digitos(azar.entero("grupo", 1, 999_999), 6), ""),
            azar("congrupo") < 0.3,
        ),
        "GRUPO_MAZ": _o_nulo(
            pc.binary_join_element_wise("M", _digitos(azar.entero("maz", 1, 9_999), 4), ""),
            azar("conmaz") < 0.1,
        ),
        "CLAVE_SPEI": pc.binary_join_element_wise("000", _digitos(c, 15), ""),
        "PAGOS_CLIENTE": _numero(estado.pagos),
        "MONTO_PAGOS": _pesos(estado.monto_pagos_centavos),
        "GESTORES": codigo("GESTOR", "gestor", 80, 3),
        "FOLIO_PLAN": _o_nulo(
            pc.binary_join_element_wise("PL", _digitos(c % 10**8, 8), ""), con_plan
        ),
        "SEGMENTO_GENERACION": _o_nulo(_texto(["GEN-A"] * n), con_plan),
        "ESTATUS_PLAN": _o_nulo(
            _de_lista(("VIGENTE", "CUMPLIDO", "CANCELADO"), azar.indice("estatusplan", 3)),
            con_plan,
        ),
        "SEMANAS_ATRASO": _numero(semanas),
        "ATRASO": _de_lista(tuple(etiquetas), tramo),
        "GENERACION_PLAN": _o_nulo(_fechas(estado.asignacion), con_plan),
        "CANCELACION_CUMPLIMIENTO_PLAN": nulos,
        "ULTIMO_ESTATUS": _o_nulo(gestion, gestionada),
        "EMPLEADO": _o_nulo(
            pc.binary_join_element_wise("E", _digitos(azar.entero("empleado", 1, 99_999), 5), ""),
            azar("conempleado") < 0.2,
        ),
        "CANAL": _de_lista(CANALES, canal),
        "ABONO_SEMANAL": _o_nulo(_pesos(abono), con_plan),
        "PLAZO": _o_nulo(_numero(plazo), con_plan),
        "MONTO_ABONADO": _o_nulo(_pesos(abono * recibidos), con_plan),
        "MONTO_PLAN": _o_nulo(_pesos(monto_plan), con_plan),
        "ENGANCHE": _o_nulo(_pesos(enganche), con_plan),
        "PAGOS_RECIBIDOS": _o_nulo(_numero(recibidos), con_plan),
        "SALDO_ANTES_DEL_PLAN": _o_nulo(_pesos(total), con_plan),
        "SALDO_ATRASADO_ANTES_PLAN": _o_nulo(_pesos(atrasado), con_plan),
        "MORATORIOS_ANTES_PLAN": _o_nulo(_pesos(moratorios), con_plan),
        "ESTATUS_PROMESA_PAGO": _o_nulo(
            _de_lista(("VIGENTE", "CUMPLIDA", "INCUMPLIDA"), azar.indice("estatuspromesa", 3)),
            con_promesa,
        ),
        "MONTO_PROMESA_PAGO": _o_nulo(_pesos(np.maximum(atrasado // 2, 5000)), con_promesa),
    }
    for k in range(1, 5):
        columnas[f"TELEFONO{k}"] = _o_nulo(telefono(f"telefono{k}"), telefonos >= k)
        tipo = _elegir(azar(f"tipotel{k}"), [0.55, 0.2, 0.15, 0.1])
        columnas[f"TIPOTEL{k}"] = _o_nulo(_de_lista(TIPOS_TELEFONO, tipo), telefonos >= k)
    return pa.table({nombre: columnas[nombre] for nombre in CONTRATO_V2.nombres})


def tabla_carrier(cartera: pa.Table) -> pa.Table:
    """CARRIER, derivada de CARTERA: una fila por cada telefono distinto de cada cliente (sus
    TELEFONO1 a TELEFONO4 y el de su aval), con las demas columnas de su fila y el telefono en
    TELEFONO. No agrega ningun cliente ni ningun telefono que CARTERA no traiga."""
    fijas = [c for c in COLUMNAS_CARRIER if c != "TELEFONO"]
    partes = []
    for orden, columna in enumerate(
        ("TELEFONO1", "TELEFONO2", "TELEFONO3", "TELEFONO4", "TEL_AVAL")
    ):
        presente = pc.is_valid(cartera[columna])
        parte = (
            cartera.filter(presente)
            .select(fijas)
            .append_column("TELEFONO", cartera[columna].filter(presente))
        )
        posicion = pc.add(
            pc.multiply(
                pa.array(np.flatnonzero(presente.to_numpy(zero_copy_only=False)), pa.int64()), 5
            ),
            orden,
        )
        partes.append(parte.append_column("_orden", posicion))
    largo = pa.concat_tables(partes).sort_by("_orden").drop_columns(["_orden"])
    tabla = largo.to_pandas(types_mapper=pd.ArrowDtype).drop_duplicates(
        ["CLIENTE_UNICO", "TELEFONO"]
    )
    return pa.Table.from_pandas(tabla, preserve_index=False).select(list(COLUMNAS_CARRIER))


# --- defectos a proposito -------------------------------------------------------------------------

DEFECTOS_V2: tuple[tuple[str, object], ...] = (
    ("SALDO_TOTAL", "-150.00"),
    ("SALDO_TOTAL", "N/D"),
    ("DIAS_ATRASO", "9999"),
    ("PRODUCTO", "HIPOTECARIO"),
    ("CANAL", "CARTA"),
    ("FECHA_ASIGNACION", "30/09/2026"),
    ("CP_CTE", "7200"),
    ("POBLACION_CTE", "MUNICIPIO QUE NO EXISTE"),
    ("SALDO", "1500.505"),
)
"""Cada defecto es invalido siempre: las pruebas cuentan con eso. El ultimo de la rotacion es un
CLIENTE_UNICO repetido."""


def contaminar(cartera: pd.DataFrame, tasa: float, semilla: int) -> tuple[pd.DataFrame, list[int]]:
    """Una copia con una fraccion `tasa` de filas invalidas, cada una con un defecto en rotacion, y
    sus posiciones de menor a mayor. El ultimo turno copia el CLIENTE_UNICO de otra fila
    contaminada: un duplicado, que el contrato rechaza en sus dos copias."""
    if not 0 <= tasa <= 1:
        raise ValueError(f"La tasa de filas invalidas debe estar entre 0 y 1: {tasa}")
    rng = np.random.default_rng([semilla, 7])
    filas = sorted(rng.choice(len(cartera), size=round(len(cartera) * tasa), replace=False))
    sucia = cartera.reset_index(drop=True).astype(object)
    for i, fila in enumerate(filas):
        turno = i % (len(DEFECTOS_V2) + 1)
        if turno == len(DEFECTOS_V2):
            sucia.at[fila, "CLIENTE_UNICO"] = sucia.at[filas[i - 1], "CLIENTE_UNICO"]
        else:
            columna, defecto = DEFECTOS_V2[turno]
            sucia.at[fila, columna] = defecto
    return sucia, [int(f) for f in filas]


# --- la cartera, como DataFrame o en un archivo ---------------------------------------------------


def generar_cartera_oficial(n: int, *, semilla: int, fecha_corte: date) -> pd.DataFrame:
    """Las 93 columnas de `n` cuentas en un DataFrame de texto. Para las pruebas: para archivos
    grandes, `escribir_cartera`, que no la arma entera."""
    estado = estado_inicial(n, semilla=semilla, fecha_corte=fecha_corte)
    tabla = tabla_cartera(estado, semilla=semilla, fecha_corte=fecha_corte)
    return tabla.to_pandas(types_mapper={pa.string(): pd.StringDtype()}.get)


@dataclass(frozen=True)
class Escrito:
    """Un archivo de cartera generado."""

    ruta: Path
    filas: int
    filas_carrier: int


def bloques_cartera(
    estado: EstadoCartera,
    *,
    semilla: int,
    fecha_corte: date,
    filas_por_bloque: int = FILAS_POR_BLOQUE,
) -> Iterator[pa.Table]:
    for inicio in range(0, len(estado), filas_por_bloque):
        parte = estado.tomar(np.arange(inicio, min(inicio + filas_por_bloque, len(estado))))
        yield tabla_cartera(parte, semilla=semilla, fecha_corte=fecha_corte)


def escribir_cartera(
    estado: EstadoCartera,
    destino: str | Path,
    *,
    semilla: int,
    fecha_corte: date,
    con_carrier: bool = True,
    filas_por_bloque: int = FILAS_POR_BLOQUE,
    tabla: pd.DataFrame | None = None,
) -> Escrito:
    """Escribe CARTERA (y CARRIER, salvo con_carrier=False) por bloques. El formato sale de la
    extension: xlsx (hojas CARTERA y CARRIER, mas una LEEME), csv (solo CARTERA) o zip (CARTERA.csv
    y CARRIER.csv). Con `tabla`, escribe esas filas tal cual, para defectos a proposito."""
    ruta = Path(destino)
    formato = ruta.suffix.lower()
    if formato not in (".xlsx", ".csv", ".zip"):
        raise ValueError(f"Formato no soportado: {ruta.name!r}. Usa .xlsx, .csv o .zip.")
    ruta.parent.mkdir(parents=True, exist_ok=True)

    def carteras() -> Iterator[pa.Table]:
        if tabla is not None:
            yield pa.Table.from_pandas(tabla.astype("string"), preserve_index=False)
            return
        yield from bloques_cartera(
            estado, semilla=semilla, fecha_corte=fecha_corte, filas_por_bloque=filas_por_bloque
        )

    def carriers() -> Iterator[pa.Table]:
        for bloque in carteras():
            yield tabla_carrier(bloque)

    if formato == ".csv":
        filas = _escribir_csv(ruta.open("wb"), carteras())
        return Escrito(ruta, filas, 0)
    if formato == ".zip":
        with zipfile.ZipFile(
            ruta, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True
        ) as paquete:
            paquete.writestr(_miembro("LEEME.txt"), LEEME)
            with paquete.open(_miembro(f"{HOJA_CARTERA}.csv"), "w", force_zip64=True) as miembro:
                filas = _escribir_csv(miembro, carteras())
            filas_carrier = 0
            if con_carrier:
                with paquete.open(
                    _miembro(f"{HOJA_CARRIER}.csv"), "w", force_zip64=True
                ) as miembro:
                    filas_carrier = _escribir_csv(miembro, carriers())
        return Escrito(ruta, filas, filas_carrier)
    return _escribir_xlsx(ruta, carteras(), carriers() if con_carrier else None)


def _miembro(nombre: str) -> zipfile.ZipInfo:
    """Un miembro comprimido de un zip generado, sin la hora en que se escribio."""
    info = zipfile.ZipInfo(nombre, date_time=FECHA_EN_ZIP)
    info.compress_type = zipfile.ZIP_DEFLATED
    info.external_attr = 0o644 << 16
    return info


def _escribir_csv(sink, bloques: Iterator[pa.Table]) -> int:
    filas = 0
    escritor = None
    try:
        for bloque in bloques:
            if escritor is None:
                escritor = pcsv.CSVWriter(sink, bloque.schema)
            escritor.write_table(bloque)
            filas += bloque.num_rows
    finally:
        if escritor is not None:
            escritor.close()
        sink.close()
    return filas


TIPADAS_EN_XLSX = {Tipo.ENTERO, Tipo.IMPORTE, Tipo.DECIMAL, Tipo.FECHA, Tipo.FECHA_HORA}
"""En un xlsx, estas columnas van como numeros y fechas de Excel, como las escribe un sistema que
exporta a Excel; las demas, como texto, para no perder los ceros a la izquierda."""


def _escribir_xlsx(ruta: Path, carteras: Iterator[pa.Table], carriers) -> Escrito:
    import openpyxl

    libro = openpyxl.Workbook(write_only=True)
    leeme = libro.create_sheet("LEEME")
    leeme.append(["Nota"])
    leeme.append([LEEME])
    filas = _hoja_xlsx(libro.create_sheet(HOJA_CARTERA), carteras, CONTRATO_V2.nombres)
    filas_carrier = 0
    if carriers is not None:
        filas_carrier = _hoja_xlsx(libro.create_sheet(HOJA_CARRIER), carriers, COLUMNAS_CARRIER)
    libro.save(ruta)
    return Escrito(ruta, filas, filas_carrier)


def _hoja_xlsx(
    hoja,
    bloques: Iterator[pa.Table],
    encabezado: tuple[str, ...],
    contrato: ContratoFuente = CONTRATO_V2,
) -> int:
    tipos = {c.nombre: c.tipo for c in contrato.columnas}
    hoja.append(list(encabezado))
    filas = 0
    for bloque in bloques:
        columnas = []
        for nombre in encabezado:
            valores = bloque[nombre].to_pylist()
            tipo = tipos.get(nombre)
            if tipo in TIPADAS_EN_XLSX:
                valores = [_celda(v, tipo) for v in valores]
            columnas.append(valores)
        for fila in zip(*columnas, strict=True):
            hoja.append(fila)
        filas += bloque.num_rows
    return filas


def _celda(valor: str | None, tipo: Tipo):
    """Un valor como lo guardaria Excel: un numero o una fecha. Si no se puede (un defecto a
    proposito), como el texto que es."""
    if valor is None:
        return None
    try:
        if tipo == Tipo.ENTERO:
            return int(valor)
        if tipo in (Tipo.IMPORTE, Tipo.DECIMAL):
            return float(valor)
        if tipo == Tipo.FECHA_HORA:
            return datetime.fromisoformat(valor.replace(" ", "T"))
        return date.fromisoformat(valor)
    except ValueError:
        return valor


# --- PAGOS: los movimientos economicos de un periodo ----------------------------------------------

CONCEPTOS = ("PAGO NORMAL", "ABONO", "LIQUIDACION")
PESO_CONCEPTO = (0.70, 0.27, 0.03)
# Comisiones genericas (5%, 8% y 10%), en puntos base: no son las de ningun contrato.
COMISIONES_PUNTOS_BASE = (500, 800, 1000)


@dataclass(frozen=True)
class Movimientos:
    """Los pagos de un periodo: las 23 columnas y, para evolucionar la cartera, de que cuenta es
    cada movimiento distinto, cuando se recibio y cuanto."""

    tabla: pa.Table
    """Las 23 columnas, como texto, con los repetidos exactos, ordenadas por Fecha_Recepción."""
    cuenta: np.ndarray
    """La posicion en el estado de la cuenta de cada movimiento distinto, sin los repetidos."""
    recepcion: np.ndarray
    """Cuando se recibio cada movimiento distinto (datetime64[s])."""
    recuperado_centavos: np.ndarray
    repetidos: int
    """Cuantas filas de la tabla son copias exactas de otra: el mismo pago reportado dos veces."""
    ajustes: int


def tabla_pagos(
    estado: EstadoCartera,
    *,
    semilla: int,
    desde: date,
    hasta: date,
    fraccion: float = 0.35,
    repetidos: float = 0.005,
    ajustes: float = 0.003,
) -> pa.Table:
    """Los movimientos de un periodo, de `desde` a `hasta` inclusive, de las cuentas de `estado`:
    las 23 columnas de `movimientos_del_periodo`."""
    return movimientos_del_periodo(
        estado,
        semilla=semilla,
        desde=desde,
        hasta=hasta,
        fraccion=fraccion,
        repetidos=repetidos,
        ajustes=ajustes,
    ).tabla


def movimientos_del_periodo(
    estado: EstadoCartera,
    *,
    semilla: int,
    desde: date,
    hasta: date,
    fraccion: float = 0.35,
    repetidos: float = 0.005,
    ajustes: float = 0.003,
) -> Movimientos:
    """Los movimientos de un periodo, de `desde` a `hasta` inclusive, de las cuentas de `estado`.

    Paga una fraccion de las cuentas, de una a cuatro veces cada una, siempre dentro del periodo, y
    con lo que el acreedor sabia de la cuenta (su territorio, su gestor, su atraso). Una liquidacion
    paga el saldo entero. Una fraccion pequena de movimientos se repite exacta, con el mismo
    cliente, el mismo segundo y el mismo importe (la llave historica de deduplicacion), y otra es un
    ajuste negativo: v0.6 no deduplica ni interpreta signos, y las pruebas lo verifican. Las 23
    columnas, como texto y en el orden del contrato, ordenadas por Fecha_Recepción.
    """
    if hasta < desde:
        raise ValueError(f"El periodo termina antes de empezar: {desde} a {hasta}.")
    rng = np.random.default_rng([semilla, desde.toordinal(), hasta.toordinal(), 3])
    pagan = np.flatnonzero(rng.random(len(estado)) < fraccion)
    veces = 1 + np.minimum(rng.poisson(0.6, size=len(pagan)), 3)
    indices = np.repeat(pagan, veces)
    cuentas = estado.tomar(indices)
    m = len(cuentas)
    c = cuentas.cliente
    azar = _Azar(c, semilla)
    geo = _geografia()
    cve_ent = geo.cve_entidad[_elegir(azar("municipio"), list(geo.peso))]

    segundos = ((hasta - desde).days + 1) * 86_400
    inicio = np.datetime64(desde, "s")
    recepcion = inicio + rng.integers(0, segundos, size=m).astype("timedelta64[s]")
    gestion = recepcion - rng.integers(0, 10 * 86_400, size=m).astype("timedelta64[s]")
    con_gestion = rng.random(m) < 0.8
    dias = cuentas.dias_atraso
    etiquetas = [e for e, _, _ in TRAMOS_ATRASO]
    tramo = np.zeros(m, dtype=np.int64)
    for j, (_, desde_tramo, hasta_tramo) in enumerate(TRAMOS_ATRASO):
        tramo[(dias >= desde_tramo) & (dias <= (10**9 if hasta_tramo is None else hasta_tramo))] = j

    pago_normal = np.maximum(cuentas.saldo_centavos // 52, 5000)
    recuperado = np.maximum((pago_normal * (0.5 + 2.5 * rng.random(m))).astype(np.int64), 5000)
    concepto = rng.choice(len(CONCEPTOS), size=m, p=PESO_CONCEPTO)
    # Una liquidacion paga el saldo entero: la cuenta sale de la cartera en el corte siguiente.
    liquida = concepto == CONCEPTOS.index("LIQUIDACION")
    recuperado = np.where(liquida, np.maximum(cuentas.saldo_centavos, 5000), recuperado)
    ajuste = rng.random(m) < ajustes
    recuperado = np.where(ajuste, -np.minimum(recuperado, 50_000), recuperado)
    cargos = np.where(rng.random(m) < 0.05, rng.integers(1_000, 5_001, size=m), 0)
    cobranza = recuperado + cargos
    plan_comision = azar.indice("comision", len(COMISIONES_PUNTOS_BASE))
    puntos = np.array(COMISIONES_PUNTOS_BASE)[plan_comision]
    comision = np.round(cobranza * puntos / 10_000).astype(np.int64)

    iso = pd.DatetimeIndex(recepcion).isocalendar()
    columnas = {
        "Año": _numero(iso["year"].to_numpy()),
        "Semana": _numero(iso["week"].to_numpy()),
        "Territorio": pc.binary_join_element_wise("TERRITORIO", _texto(list(cve_ent)), " "),
        "Zona": pc.binary_join_element_wise("ZONA", _digitos(azar.entero("zona", 1, 12), 2), " "),
        "Cliente_Unico": pc.binary_join_element_wise("CU", _digitos(c, 10), ""),
        "Fecha_Recepción": _instantes(recepcion),
        "Segmento": _de_lista(tuple(etiquetas), tramo),
        "Gerencia": pc.binary_join_element_wise(
            "GERENCIA", _digitos(azar.entero("gerencia", 1, 6), 2), " "
        ),
        "Tipo_Cartera": _de_lista(
            ("TRADICIONAL", "ESPECIAL"), (azar("tipo_cartera") >= 0.9).astype(np.int64)
        ),
        "Producto": _de_lista(PRODUCTOS, _producto(c, semilla)),
        "Campaña": pc.binary_join_element_wise(
            "CAMP", _digitos(azar.entero("campania", 1, 6), 2), "-"
        ),
        "Gestor": pc.binary_join_element_wise(
            "GESTOR", _digitos(azar.entero("gestor", 1, 80), 3), " "
        ),
        "Días_de_Atraso": _numero(dias),
        "Semanas_de_Atraso": _numero(-(-dias // 7)),
        "Plan_de_Pago": _de_lista(("N", "S"), (azar("plan") < 0.15).astype(np.int64)),
        "Fecha_de_Gestion": _o_nulo(_instantes(gestion), con_gestion),
        "Recuperación_por_Gestión": _pesos(recuperado),
        "Concepto_Cálculo": pc.if_else(pa.array(ajuste), "AJUSTE", _de_lista(CONCEPTOS, concepto)),
        "Cargos_Automáticos": _pesos(cargos),
        "Captación": _pesos(recuperado),
        "Cobranza_Total": _pesos(cobranza),
        "Porcentaje_Comision": _de_lista(
            tuple(str(p / 10_000) for p in COMISIONES_PUNTOS_BASE), plan_comision
        ),
        "Monto_Comision": _pesos(comision),
    }
    tabla = pa.table({nombre: columnas[nombre] for nombre in CONTRATO_PAGOS.nombres})
    # Repetidos exactos: mismo cliente, mismo segundo, mismo importe. Se conservan los dos.
    copias = rng.choice(m, size=round(m * repetidos), replace=False) if m else np.array([], int)
    tabla = pa.concat_tables([tabla, tabla.take(pa.array(np.sort(copias), type=pa.int64()))])
    orden = pc.sort_indices(tabla, sort_keys=[("Fecha_Recepción", "ascending")])
    return Movimientos(
        tabla=tabla.take(orden),
        cuenta=indices,
        recepcion=recepcion,
        recuperado_centavos=recuperado,
        repetidos=len(copias),
        ajustes=int(ajuste.sum()),
    )


def _instantes(segundos: np.ndarray) -> pa.Array:
    """AAAA-MM-DD HH:MM:SS, como lo exporta un sistema de pagos."""
    texto = pc.cast(pa.array(segundos.astype("datetime64[s]")), pa.string())
    return pc.replace_substring(texto, "T", " ")


HOJA_PAGOS = "PAGOS"


@dataclass(frozen=True)
class EscritoPagos:
    """Un archivo de pagos generado."""

    ruta: Path
    filas: int


def escribir_pagos(tabla: pa.Table | pd.DataFrame, destino: str | Path) -> EscritoPagos:
    """Escribe los movimientos en un archivo; el formato sale de la extension: csv, zip (con
    PAGOS.csv) o xlsx (hoja PAGOS, con numeros y fechas de Excel donde el contrato los tipa). Con un
    DataFrame, escribe esas filas tal cual, para defectos a proposito."""
    if isinstance(tabla, pd.DataFrame):
        tabla = pa.Table.from_pandas(tabla.astype("string"), preserve_index=False)
    ruta = Path(destino)
    formato = ruta.suffix.lower()
    if formato not in (".xlsx", ".csv", ".zip"):
        raise ValueError(f"Formato no soportado: {ruta.name!r}. Usa .xlsx, .csv o .zip.")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    if formato == ".csv":
        return EscritoPagos(ruta, _escribir_csv(ruta.open("wb"), iter([tabla])))
    if formato == ".zip":
        with zipfile.ZipFile(
            ruta, "w", compression=zipfile.ZIP_DEFLATED, allowZip64=True
        ) as paquete:
            paquete.writestr(_miembro("LEEME.txt"), LEEME)
            with paquete.open(_miembro(f"{HOJA_PAGOS}.csv"), "w", force_zip64=True) as miembro:
                filas = _escribir_csv(miembro, iter([tabla]))
        return EscritoPagos(ruta, filas)
    import openpyxl

    libro = openpyxl.Workbook(write_only=True)
    filas = _hoja_xlsx(
        libro.create_sheet(HOJA_PAGOS), iter([tabla]), CONTRATO_PAGOS.nombres, CONTRATO_PAGOS
    )
    libro.save(ruta)
    return EscritoPagos(ruta, filas)


# --- el escenario longitudinal: varios cortes y los pagos entre ellos -----------------------------


@dataclass(frozen=True)
class Evolucion:
    """Que paso con las cuentas de un corte al siguiente."""

    continuan: int
    liquidadas: int
    """Lo que pagaron en el periodo cubrio su saldo: salen de la cartera."""
    retiradas: int
    """El acreedor las retiro del despacho."""
    altas: int
    """Cuentas nuevas, de clientes que nunca habian estado en la cartera."""


def evolucionar(
    estado: EstadoCartera,
    movimientos: Movimientos,
    *,
    semilla: int,
    corte: date,
    siguiente: date,
    usados: set[int],
    tasa_altas: float = 0.02,
    tasa_retiros: float = 0.01,
) -> tuple[EstadoCartera, Evolucion]:
    """Las cuentas del corte `siguiente`, a partir de las de `corte` y de sus pagos del periodo.

    - Cada cuenta resta de su SALDO lo que recupero en el periodo, contando una vez cada movimiento
      (un repetido exacto es el mismo pago reportado dos veces); un ajuste negativo lo aumenta. Si
      llega a cero o menos, se liquido y sale de la cartera.
    - El acreedor retira al azar una fraccion `tasa_retiros` de las demas.
    - Las que continuan conservan su identidad y sus atributos fijos. Su atraso vuelve a 0 si
      pagaron algo en el periodo; si no, crece con los dias del periodo, hasta el tope. Sus pagos
      positivos se suman a su historial y el ultimo es el de su ultimo pago.
    - Llegan cuentas nuevas, una fraccion `tasa_altas` del corte, de clientes que nunca habian
      estado en la cartera (`usados`), asignadas dentro del periodo.
    """
    dias_periodo = (siguiente - corte).days
    if dias_periodo < 1:
        raise ValueError(f"El corte siguiente tiene que ser posterior: {corte} y {siguiente}.")
    n = len(estado)
    cuenta, centavos = movimientos.cuenta, movimientos.recuperado_centavos
    pagado = np.zeros(n, dtype=np.int64)
    np.add.at(pagado, cuenta, centavos)
    positivo = centavos > 0
    veces = np.bincount(cuenta[positivo], minlength=n)
    monto = np.zeros(n, dtype=np.int64)
    np.add.at(monto, cuenta[positivo], centavos[positivo])
    # El ultimo pago positivo de cada cuenta: el de recepcion mas reciente.
    orden = np.lexsort((movimientos.recepcion[positivo], cuenta[positivo]))
    de_cuenta = cuenta[positivo][orden]
    ultimo = np.r_[de_cuenta[1:] != de_cuenta[:-1], True] if len(de_cuenta) else de_cuenta > 0
    con_pago = de_cuenta[ultimo]

    rng = np.random.default_rng([semilla, siguiente.toordinal(), 5])
    saldo = estado.saldo_centavos - pagado
    liquidadas = saldo <= 0
    retiradas = ~liquidadas & (rng.random(n) < tasa_retiros)
    continuan = ~(liquidadas | retiradas)
    dias = np.where(pagado > 0, 0, np.minimum(estado.dias_atraso + dias_periodo, TOPE_ATRASO))
    ultimo_pago = estado.ultimo_pago.copy()
    ultimo_pago[con_pago] = movimientos.recepcion[positivo][orden][ultimo].astype("datetime64[D]")
    ultimo_centavos = estado.ultimo_pago_centavos.copy()
    ultimo_centavos[con_pago] = centavos[positivo][orden][ultimo]
    siguen = EstadoCartera(
        cliente=estado.cliente,
        saldo_centavos=saldo,
        moratorios_centavos=(np.maximum(saldo, 0) * 0.0015 * dias).astype(np.int64),
        dias_atraso=dias,
        atraso_maximo=np.maximum(estado.atraso_maximo, dias),
        asignacion=estado.asignacion,
        pagos=estado.pagos + veces,
        monto_pagos_centavos=estado.monto_pagos_centavos + monto,
        ultimo_pago=ultimo_pago,
        ultimo_pago_centavos=ultimo_centavos,
    ).tomar(np.flatnonzero(continuan))

    altas = int(rng.binomial(n, tasa_altas)) if n else 0
    nuevas = _cuentas_nuevas(nuevos_clientes(altas, rng, usados), rng, semilla, siguiente)
    llegada = np.datetime64(siguiente, "D") - rng.integers(0, dias_periodo, size=altas).astype(
        "timedelta64[D]"
    )
    nuevas = replace(nuevas, asignacion=llegada.astype("datetime64[D]"))
    juntas = EstadoCartera(
        **{
            campo: np.concatenate([getattr(siguen, campo), getattr(nuevas, campo)])
            for campo in EstadoCartera.__dataclass_fields__
        }
    )
    evolucion = Evolucion(
        continuan=int(continuan.sum()),
        liquidadas=int(liquidadas.sum()),
        retiradas=int(retiradas.sum()),
        altas=altas,
    )
    return juntas, evolucion


INVARIANTES = (
    "Determinista: la misma semilla y los mismos parametros producen los mismos archivos, byte por "
    "byte, y el mismo manifiesto.",
    "Cada corte cumple cartera/v2 y cada periodo cumple pagos/v1; CLIENTE_UNICO es unico en cada "
    "corte.",
    "Una alta nunca reusa un cliente que ya estuvo en la cartera, aunque haya salido.",
    "Una cuenta que continua conserva su identidad y sus atributos fijos: nombre, domicilio, "
    "telefonos, producto, municipio y fecha de asignacion.",
    "Los pagos de un periodo son de cuentas del corte que lo abre y se reciben dentro del periodo: "
    "despues de ese corte y hasta el corte que lo cierra, inclusive.",
    "Para cada cuenta que continua, su SALDO es el anterior menos lo que recupero en el periodo, "
    "contando una vez cada movimiento: un repetido exacto es el mismo pago reportado dos veces.",
    "Una cuenta cuyos pagos del periodo cubren su saldo se liquida y sale de la cartera.",
    "El atraso de una cuenta que continua vuelve a 0 si pago algo en el periodo; si no, crece con "
    "los dias del periodo, hasta el tope.",
    "Las cuentas de un corte son las que continuan mas las altas: las demas se liquidaron o se "
    "retiraron, y el manifiesto lo cuenta.",
    "CARRIER no introduce clientes ni telefonos que su corte no traiga, y todas las cuentas son de "
    "un solo despacho.",
)
"""Lo que cumple todo escenario. Las pruebas lo verifican sobre los archivos escritos."""

INVARIANTES_LIFECYCLE = (
    "El lifecycle es sintetico y determinista: la misma semilla da los mismos eventos, y "
    "agregarlo no cambia ningun corte ni ningun pago.",
    "Cada gestion es de una cuenta del corte que abre su periodo y cumple lifecycle/v1; cada "
    "promesa, convenio, cancelacion y anulacion se refiere por su llave a un evento de su mismo "
    "archivo.",
    "Ningun evento ocurre despues del ultimo corte, y el escenario termina antes del dia en que se "
    "genera: un evento operacional no ocurre en el futuro.",
    "Un grupo tardio (una gestion con lo que nacio de ella) ocurrio en un periodo y llega en el "
    "archivo del siguiente.",
)
"""Lo que cumple el lifecycle de un escenario que lo trae."""

MANIFIESTO = "escenario.json"


@dataclass(frozen=True)
class Escenario:
    """Un escenario escrito: su directorio y su manifiesto."""

    destino: Path
    manifiesto: dict


def generar_escenario(
    destino: str | Path,
    *,
    cuentas: int,
    cortes: int,
    primer_corte: date,
    semilla: int,
    dias_entre_cortes: int = 7,
    formato: str = "zip",
    con_carrier: bool = True,
    tasa_altas: float = 0.02,
    tasa_retiros: float = 0.01,
    lifecycle: bool = False,
    intensidad_lifecycle: float = 1.0,
    hoy: date | None = None,
) -> Escenario:
    """Escribe `cortes` cortes de la misma cartera, cada `dias_entre_cortes` dias desde
    `primer_corte`, y los pagos de cada periodo entre un corte y el siguiente; y un manifiesto,
    escenario.json, con cada archivo, su SHA-256, sus conteos y las invariantes que cumple.

    Cada corte sale del anterior con `evolucionar`: los pagos del periodo bajan los saldos y curan
    el atraso, las cuentas liquidadas y las retiradas salen y llegan altas. Todo se deriva de la
    semilla: el mismo escenario sale igual, byte por byte, en csv y en zip.

    Con `lifecycle`, tambien los eventos operacionales de cada periodo (`generador.lifecycle`), en
    un JSONL comprimido por periodo. Exige que el escenario termine antes de `hoy`: un evento
    operacional no ocurre en el futuro.
    """
    if cortes < 1:
        raise ValueError("Un escenario tiene al menos un corte.")
    if dias_entre_cortes < 1:
        raise ValueError("Entre un corte y el siguiente pasa al menos un dia.")
    extension = formato.lower().lstrip(".")
    if extension not in ("xlsx", "csv", "zip"):
        raise ValueError(f"Formato no soportado: {formato!r}. Usa xlsx, csv o zip.")
    ultimo_corte = primer_corte + timedelta(days=dias_entre_cortes * (cortes - 1))
    if lifecycle and ultimo_corte >= (hoy or date.today()):
        raise ValueError(
            f"Con el lifecycle el escenario tiene que terminar antes de hoy, y su ultimo corte es "
            f"el {ultimo_corte.isoformat()}: un evento operacional no ocurre en el futuro. Usa un "
            "primer corte anterior."
        )
    directorio = Path(destino)
    directorio.mkdir(parents=True, exist_ok=True)
    periodos = None
    if lifecycle:
        from motor_cartera.generador.lifecycle import Periodos

        periodos = Periodos(directorio)
    usados: set[int] = set()
    estado = estado_inicial(cuentas, semilla=semilla, fecha_corte=primer_corte, usados=usados)
    registro_cortes: list[dict] = []
    registro_periodos: list[dict] = []
    evolucion: Evolucion | None = None
    corte = primer_corte
    for numero in range(cortes):
        ruta = directorio / f"cartera_oficial_{corte.isoformat()}.{extension}"
        escrito = escribir_cartera(
            estado, ruta, semilla=semilla, fecha_corte=corte, con_carrier=con_carrier
        )
        registro_cortes.append(
            {
                "fecha_corte": corte.isoformat(),
                "archivo": ruta.name,
                "sha256": _sha256(ruta),
                "cuentas": escrito.filas,
                "filas_carrier": escrito.filas_carrier,
                **(asdict(evolucion) if evolucion is not None else {}),
            }
        )
        if numero == cortes - 1:
            break
        siguiente = corte + timedelta(days=dias_entre_cortes)
        desde = corte + timedelta(days=1)
        movimientos = movimientos_del_periodo(estado, semilla=semilla, desde=desde, hasta=siguiente)
        ruta = directorio / f"pagos_oficial_{desde.isoformat()}_{siguiente.isoformat()}.{extension}"
        escritos = escribir_pagos(movimientos.tabla, ruta)
        registro_periodos.append(
            {
                "desde": desde.isoformat(),
                "hasta": siguiente.isoformat(),
                "archivo": ruta.name,
                "sha256": _sha256(ruta),
                "movimientos": escritos.filas,
                "repetidos_exactos": movimientos.repetidos,
                "ajustes": movimientos.ajustes,
            }
        )
        if periodos is not None:
            from motor_cartera.generador.lifecycle import grupos_del_periodo

            periodos.agregar(
                desde,
                siguiente,
                grupos_del_periodo(
                    estado,
                    movimientos,
                    semilla=semilla,
                    desde=desde,
                    hasta=siguiente,
                    intensidad=intensidad_lifecycle,
                    con_tardios=numero < cortes - 2,
                ),
            )
        estado, evolucion = evolucionar(
            estado,
            movimientos,
            semilla=semilla,
            corte=corte,
            siguiente=siguiente,
            usados=usados,
            tasa_altas=tasa_altas,
            tasa_retiros=tasa_retiros,
        )
        corte = siguiente
    manifiesto = {
        "version": "escenario/v1",
        "semilla": semilla,
        "cuentas_iniciales": cuentas,
        "dias_entre_cortes": dias_entre_cortes,
        "tasa_altas": tasa_altas,
        "tasa_retiros": tasa_retiros,
        "formato": extension,
        "contratos": {"cortes": CONTRATO_V2.version, "periodos": CONTRATO_PAGOS.version},
        "cortes": registro_cortes,
        "periodos": registro_periodos,
        "invariantes": list(INVARIANTES),
    }
    if periodos is not None:
        manifiesto["contratos"]["lifecycle"] = "lifecycle/v1"
        manifiesto["intensidad_lifecycle"] = intensidad_lifecycle
        manifiesto["lifecycle"] = [archivo.manifiesto() for archivo in periodos.archivos]
        manifiesto["invariantes"] += list(INVARIANTES_LIFECYCLE)
    (directorio / MANIFIESTO).write_text(
        json.dumps(manifiesto, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return Escenario(directorio, manifiesto)


def _sha256(ruta: Path) -> str:
    digesto = hashlib.sha256()
    with ruta.open("rb") as archivo:
        while bloque := archivo.read(1 << 20):
            digesto.update(bloque)
    return digesto.hexdigest()
