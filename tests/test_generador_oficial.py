"""El generador de las fuentes oficiales, sin base de datos: sintetico, reproducible, coherente y
siempre dentro del contrato que simula."""

from __future__ import annotations

import hashlib
import io
import zipfile
from datetime import date

import numpy as np
import openpyxl
import pandas as pd
import pyarrow as pa
import pytest

from motor_cartera.contratos.cartera_v2 import COLUMNAS_CARRIER, CONTRATO_V2
from motor_cartera.contratos.fuente import juzgar_lote
from motor_cartera.fuentes.proyeccion import rechazos_geograficos
from motor_cartera.generador.oficial import (
    PERFILES,
    contaminar,
    escribir_cartera,
    estado_inicial,
    generar_cartera_oficial,
    tabla_carrier,
    tabla_cartera,
)

CORTE = date(2026, 9, 30)
TELEFONOS = ("TELEFONO1", "TELEFONO2", "TELEFONO3", "TELEFONO4", "TEL_AVAL")


def _texto(df: pd.DataFrame) -> pd.DataFrame:
    return df.astype("string")


def test_los_perfiles_de_escala():
    assert PERFILES == {
        "XS": 1_000,
        "S": 10_000,
        "M": 100_000,
        "L": 250_000,
        "XL": 500_000,
        "XXL": 1_000_000,
    }


@pytest.mark.parametrize("semilla", [1, 31416, 2026])
def test_todo_lo_que_genera_cumple_el_contrato_y_se_proyecta(semilla):
    cartera = generar_cartera_oficial(3_000, semilla=semilla, fecha_corte=CORTE)

    juicio = juzgar_lote(_texto(cartera), CONTRATO_V2)

    assert juicio.rechazos == {}
    assert rechazos_geograficos(juicio.canonicos) == {}
    assert cartera["CLIENTE_UNICO"].is_unique
    assert list(cartera.columns) == list(CONTRATO_V2.nombres)


def test_es_reproducible_con_la_misma_semilla():
    una = generar_cartera_oficial(500, semilla=7, fecha_corte=CORTE)
    otra = generar_cartera_oficial(500, semilla=7, fecha_corte=CORTE)
    distinta = generar_cartera_oficial(500, semilla=8, fecha_corte=CORTE)

    pd.testing.assert_frame_equal(una, otra)
    assert not una.equals(distinta)


def test_lo_fijo_de_un_cliente_no_depende_del_lote_en_que_se_genera():
    # Nombre, domicilio, telefonos y municipio salen del cliente: igual en un lote de 2,000 que en
    # uno de tres, y en cualquier orden.
    estado = estado_inicial(2_000, semilla=5, fecha_corte=CORTE)
    completa = tabla_cartera(estado, semilla=5, fecha_corte=CORTE).to_pandas()
    indices = np.array([1999, 7, 1000])
    parcial = tabla_cartera(estado.tomar(indices), semilla=5, fecha_corte=CORTE).to_pandas()

    pd.testing.assert_frame_equal(
        parcial.reset_index(drop=True), completa.iloc[indices].reset_index(drop=True)
    )


def test_los_datos_sinteticos_no_pueden_ser_de_nadie():
    cartera = generar_cartera_oficial(2_000, semilla=3, fecha_corte=CORTE)

    for columna in TELEFONOS:
        telefonos = cartera[columna].dropna()
        # Ningun numero nacional de Mexico empieza con 0: no se le puede marcar a nadie.
        assert telefonos.str.fullmatch(r"0\d{9}").all()
    assert cartera["CLAVE_SPEI"].str.fullmatch(r"000\d{15}").all()  # 000 no es ningun banco
    assert cartera["LATITUD"].isna().all() and cartera["LONGITUD"].isna().all()


def test_lo_que_genera_es_coherente():
    cartera = generar_cartera_oficial(2_000, semilla=4, fecha_corte=CORTE)
    centavos = {
        c: (cartera[c].astype(float) * 100).round().astype(int)
        for c in ("SALDO", "MORATORIOS", "SALDO_TOTAL")
    }
    dias = cartera["DIAS_ATRASO"].astype(int)

    assert (centavos["SALDO"] + centavos["MORATORIOS"] == centavos["SALDO_TOTAL"]).all()
    assert (cartera["SEMANAS_ATRASO"].astype(int) == -(-dias // 7)).all()
    assert (cartera["ATRASO_MAXIMO"].astype(int) >= dias).all()
    for k in range(1, 5):
        assert (cartera[f"TELEFONO{k}"].isna() == cartera[f"TIPOTEL{k}"].isna()).all()
    # Telefonos en orden: si tiene el tercero, tiene el primero y el segundo.
    assert (cartera["TELEFONO3"].isna() | cartera["TELEFONO2"].notna()).all()
    pago = cartera["FECHA_ULTIMO_PAGO"].notna()
    assert (pago == cartera["IMP_ULTIMO_PAGO"].notna()).all()
    assert (pd.to_datetime(cartera.loc[pago, "FECHA_ULTIMO_PAGO"]) < pd.Timestamp(CORTE)).all()


def test_se_concentra_como_una_cartera_regional():
    cartera = generar_cartera_oficial(10_000, semilla=31416, fecha_corte=CORTE)

    assert (cartera["ESTADO_CTE"] == "PUEBLA").mean() == pytest.approx(0.70, abs=0.02)
    poblaciones = cartera.loc[cartera["ESTADO_CTE"] == "PUEBLA", "POBLACION_CTE"].value_counts()
    assert poblaciones.index[0] == "PUEBLA"  # la capital, por su poblacion del censo
    assert poblaciones.iloc[0] > 10 * poblaciones.median()


def test_carrier_sale_de_los_telefonos_de_cartera_y_de_nada_mas():
    estado = estado_inicial(1_500, semilla=6, fecha_corte=CORTE)
    cartera = tabla_cartera(estado, semilla=6, fecha_corte=CORTE)

    carrier = tabla_carrier(cartera).to_pandas()
    datos = cartera.to_pandas()

    assert list(carrier.columns) == list(COLUMNAS_CARRIER)
    telefonos = {
        (c, t)
        for col in TELEFONOS
        for c, t in datos[["CLIENTE_UNICO", col]].dropna().itertuples(index=False)
    }
    assert (
        set(carrier[["CLIENTE_UNICO", "TELEFONO"]].itertuples(index=False, name=None)) == telefonos
    )
    assert len(carrier) == len(telefonos)
    # Cada fila de CARRIER es la de su cliente en CARTERA, salvo los telefonos.
    uno = carrier.iloc[0]
    fila = datos.set_index("CLIENTE_UNICO").loc[uno["CLIENTE_UNICO"]]
    for columna in COLUMNAS_CARRIER:
        if columna not in ("CLIENTE_UNICO", "TELEFONO"):
            assert _igual(uno[columna], fila[columna]), columna


def _igual(a, b) -> bool:
    if pd.isna(a) or pd.isna(b):
        return pd.isna(a) and pd.isna(b)
    return a == b


def test_contaminar_marca_exactamente_lo_que_el_contrato_rechaza():
    cartera = generar_cartera_oficial(1_000, semilla=9, fecha_corte=CORTE)

    sucia, filas = contaminar(cartera, 0.05, semilla=9)

    juicio = juzgar_lote(_texto(sucia), CONTRATO_V2)
    rechazadas = set(juicio.rechazos) | set(rechazos_geograficos(juicio.canonicos))
    repetidas = set(sucia.index[sucia["CLIENTE_UNICO"].duplicated(keep=False)])
    assert len(filas) == 50
    assert sorted(rechazadas | repetidas) == filas


def test_escribe_por_bloques_lo_mismo_que_de_un_solo_golpe(tmp_path):
    estado = estado_inicial(250, semilla=10, fecha_corte=CORTE)

    uno = escribir_cartera(estado, tmp_path / "uno.csv", semilla=10, fecha_corte=CORTE)
    bloques = escribir_cartera(
        estado, tmp_path / "bloques.csv", semilla=10, fecha_corte=CORTE, filas_por_bloque=40
    )

    assert uno.filas == bloques.filas == 250
    assert uno.ruta.read_bytes() == bloques.ruta.read_bytes()


def test_cada_formato_trae_sus_hojas(tmp_path):
    estado = estado_inicial(60, semilla=11, fecha_corte=CORTE)

    xlsx = escribir_cartera(estado, tmp_path / "c.xlsx", semilla=11, fecha_corte=CORTE)
    paquete = escribir_cartera(estado, tmp_path / "c.zip", semilla=11, fecha_corte=CORTE)
    solo = escribir_cartera(
        estado, tmp_path / "s.zip", semilla=11, fecha_corte=CORTE, con_carrier=False
    )

    libro = openpyxl.load_workbook(xlsx.ruta, read_only=True)
    assert libro.sheetnames == ["LEEME", "CARTERA", "CARRIER"]
    assert xlsx.filas == 60 and xlsx.filas_carrier > 60
    with zipfile.ZipFile(paquete.ruta) as z:
        assert z.namelist() == ["LEEME.txt", "CARTERA.csv", "CARRIER.csv"]
    with zipfile.ZipFile(solo.ruta) as z:
        assert z.namelist() == ["LEEME.txt", "CARTERA.csv"]
    assert paquete.filas_carrier == xlsx.filas_carrier and solo.filas_carrier == 0


def test_en_xlsx_los_importes_y_las_fechas_van_como_numeros_y_fechas(tmp_path):
    estado = estado_inicial(5, semilla=12, fecha_corte=CORTE)
    escrito = escribir_cartera(
        estado, tmp_path / "c.xlsx", semilla=12, fecha_corte=CORTE, con_carrier=False
    )

    hoja = openpyxl.load_workbook(escrito.ruta, read_only=True)["CARTERA"]
    encabezado, fila = list(hoja.iter_rows(min_row=1, max_row=2, values_only=True))
    # openpyxl no escribe las celdas vacias del final de una fila.
    fila = list(fila) + [None] * (len(encabezado) - len(fila))
    valores = dict(zip(encabezado, fila, strict=True))
    assert isinstance(valores["SALDO_TOTAL"], float)
    assert isinstance(valores["DIAS_ATRASO"], int)
    assert valores["FECHA_ASIGNACION"].year == 2026 or valores["FECHA_ASIGNACION"].year == 2025
    assert isinstance(valores["CP_CTE"], str)  # los codigos, como texto: sin perder ceros


def test_un_formato_que_no_se_escribe(tmp_path):
    with pytest.raises(ValueError, match="Formato no soportado"):
        escribir_cartera(
            estado_inicial(3, semilla=1, fecha_corte=CORTE),
            tmp_path / "c.json",
            semilla=1,
            fecha_corte=CORTE,
        )


def test_la_tabla_de_un_bloque_es_texto_en_el_orden_oficial():
    tabla = tabla_cartera(
        estado_inicial(10, semilla=2, fecha_corte=CORTE), semilla=2, fecha_corte=CORTE
    )

    assert tabla.column_names == list(CONTRATO_V2.nombres)
    assert all(campo.type == pa.string() for campo in tabla.schema)
    assert (
        hashlib.sha256(str(tabla.to_pylist()).encode()).hexdigest()
        == hashlib.sha256(
            str(
                tabla_cartera(
                    estado_inicial(10, semilla=2, fecha_corte=CORTE), semilla=2, fecha_corte=CORTE
                ).to_pylist()
            ).encode()
        ).hexdigest()
    )


def test_el_contenido_del_zip_es_texto_csv_legible(tmp_path):
    estado = estado_inicial(30, semilla=13, fecha_corte=CORTE)
    escrito = escribir_cartera(estado, tmp_path / "c.zip", semilla=13, fecha_corte=CORTE)

    with zipfile.ZipFile(escrito.ruta) as z:
        cartera = pd.read_csv(io.BytesIO(z.read("CARTERA.csv")), dtype=str)
    assert list(cartera.columns) == list(CONTRATO_V2.nombres)
    assert len(cartera) == 30
